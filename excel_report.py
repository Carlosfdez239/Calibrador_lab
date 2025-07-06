from openpyxl import Workbook, load_workbook
from datetime import datetime
import os

def registrar_parada(info):
    archivo = "noria_registro.xlsx"
    if not os.path.exists(archivo):
        wb = Workbook()
        ws = wb.active
        ws.title = "Registro"
        ws.append([
            "Ángulo Programado (°)", "Ángulo Real (°)", "Tiempo (ms)",
            "Aceleración Inicial (°/s²)", "Aceleración de Viaje (°/s²)",
            "Aceleración Final (°/s²)", "Temperatura Driver (°C)",
            "Corriente (mA)", "Timestamp"
        ])
    else:
        wb = load_workbook(archivo)
        ws = wb.active

    ws.append([
        info["programado"], info["real"], info["tiempo_ms"],
        info["acc_ini"], info["acc_viaje"], info["acc_fin"],
        info["temperatura"], info["corriente"],
        datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    ])
    wb.save(archivo)